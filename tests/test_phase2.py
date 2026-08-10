"""Bulk actions, quick add, keepers, sell packages, exports."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from catalog.models import Printing
from catalog.normalize import normalize_name, slug_display_name
from collection.models import Holding, QuantityChange, SellPackage, SellPackageItem
from collection.services import (
    QuantityError,
    apply_bulk,
    finalize_sale,
    preview_bulk,
    set_quantity,
    undo_batch,
)
from exports import services as export_services


def reserve(card, superuser, quantity=1):
    """Reserve copies the new way: an item in a kind=KEEP package."""
    package, _ = SellPackage.objects.get_or_create(
        name="Keep", game=card.expansion.game, owner=superuser,
        kind=SellPackage.Kind.KEEP,
    )
    SellPackageItem.objects.update_or_create(
        package=package, card=card, defaults={"quantity": quantity}
    )
    return package


@pytest.fixture
def cards(expansion):
    made = []
    for slug, price in [("Card-A", "1.00"), ("Card-B", "2.00"), ("Card-C", "4.00")]:
        made.append(Printing.objects.create(
            expansion=expansion, slug=slug, display_name=slug_display_name(slug),
            name_normalized=normalize_name(slug),
            current_price_from=price, current_price_trend=price, current_price_30d=price,
        ))
    return made


@pytest.fixture
def logged_in(client, superuser):
    client.force_login(superuser)
    return client


@pytest.mark.django_db
class TestBulkServices:
    def test_preview_counts_only_real_changes(self, cards, superuser):
        set_quantity(cards[0], 1, user=superuser)
        preview = preview_bulk(cards, "set", 1, user=superuser)
        assert preview["total"] == 3
        assert preview["unchanged"] == 1  # Card-A already at 1
        assert {c.slug for c, _, _ in preview["changing"]} == {"Card-B", "Card-C"}

    def test_apply_set_writes_one_batch(self, cards, superuser):
        batch_id, changed = apply_bulk(cards, "set", 1, user=superuser)
        assert changed == 3
        assert Holding.objects.filter(quantity=1).count() == 3
        batch = QuantityChange.objects.filter(batch_id=batch_id)
        assert batch.count() == 3
        assert all(c.reason == QuantityChange.Reason.BULK_EDIT for c in batch)

    def test_apply_increment_and_zero(self, cards, superuser):
        apply_bulk(cards, "set", 2, user=superuser)
        apply_bulk(cards, "inc", 0, user=superuser)
        assert set(Holding.objects.values_list("quantity", flat=True)) == {3}
        apply_bulk(cards, "zero", 0, user=superuser)
        assert set(Holding.objects.values_list("quantity", flat=True)) == {0}

    def test_undo_batch_restores_everything(self, cards, superuser):
        set_quantity(cards[0], 5, user=superuser)
        batch_id, _ = apply_bulk(cards, "set", 1, user=superuser)
        reverted, skipped = undo_batch(batch_id, user=superuser)
        assert (reverted, skipped) == (3, 0)
        assert Holding.objects.get(card=cards[0]).quantity == 5
        assert Holding.objects.get(card=cards[1]).quantity == 0

    def test_undo_batch_skips_superseded_rows(self, cards, superuser):
        batch_id, _ = apply_bulk(cards, "set", 1, user=superuser)
        set_quantity(cards[0], 9, user=superuser)  # edited after the bulk action
        reverted, skipped = undo_batch(batch_id, user=superuser)
        assert (reverted, skipped) == (2, 1)
        assert Holding.objects.get(card=cards[0]).quantity == 9  # not clobbered

    def test_invalid_op_and_value_rejected(self, cards, superuser):
        with pytest.raises(QuantityError):
            apply_bulk(cards, "nonsense", 1, user=superuser)
        with pytest.raises(QuantityError):
            apply_bulk(cards, "set", 1000, user=superuser)


@pytest.mark.django_db
class TestBulkViews:
    def _post_data(self, game, expansion, **extra):
        data = {"game": game.code, "expansion": expansion.slug, "q": "", "owned": "0"}
        data.update(extra)
        return data

    def test_preview_then_apply_then_undo(self, logged_in, game, expansion, cards):
        preview = logged_in.post(
            "/holdings/bulk/preview/",
            self._post_data(game, expansion, op="set", value="1"),
        )
        assert preview.status_code == 200
        assert "3 cards will change" in preview.text.replace("\n", " ").replace("  ", " ") or "3</strong>" in preview.text

        applied = logged_in.post(
            "/holdings/bulk/apply/", self._post_data(game, expansion, op="set", value="1")
        )
        assert applied.status_code == 200
        assert applied.headers.get("HX-Trigger") == "qty-changed"
        assert Holding.objects.filter(quantity=1).count() == 3

        batch_id = QuantityChange.objects.first().batch_id
        undone = logged_in.post(
            f"/holdings/bulk/{batch_id}/undo/", self._post_data(game, expansion)
        )
        assert undone.status_code == 200
        assert Holding.objects.filter(quantity=1).count() == 0

    def test_scope_respects_the_visible_filter(self, logged_in, game, expansion, cards, superuser):
        set_quantity(cards[0], 1, user=superuser)
        # owned-only scope must touch only Card-A
        logged_in.post(
            "/holdings/bulk/apply/",
            self._post_data(game, expansion, op="set", value="7", owned="1"),
        )
        assert Holding.objects.get(card=cards[0]).quantity == 7
        assert not Holding.objects.filter(card=cards[1]).exists()

    def test_search_filter_scopes_bulk(self, logged_in, game, expansion, cards):
        logged_in.post(
            "/holdings/bulk/apply/",
            self._post_data(game, expansion, op="set", value="4", q="card-b"),
        )
        assert Holding.objects.get(card=cards[1]).quantity == 4
        assert Holding.objects.count() == 1

    def test_invalid_value_shows_error_dialog(self, logged_in, game, expansion, cards):
        response = logged_in.post(
            "/holdings/bulk/apply/",
            self._post_data(game, expansion, op="set", value="5000"),
        )
        assert "Cannot apply" in response.text
        assert Holding.objects.count() == 0


@pytest.mark.django_db
class TestQuickAdd:
    def test_requires_three_letters(self, logged_in, cards):
        assert "at least 3 letters" in logged_in.get("/holdings/quick-add/?q=ca").text

    def test_finds_across_games_and_sets_quantity(self, logged_in, cards):
        response = logged_in.get("/holdings/quick-add/?q=card-b")
        assert "Card B" in response.text
        logged_in.post(f"/holdings/qty/{cards[1].pk}/", {"op": "set", "value": "3"})
        assert Holding.objects.get(card=cards[1]).quantity == 3

    def test_empty_query_shows_no_results(self, logged_in, cards):
        response = logged_in.get("/holdings/quick-add/")
        assert "Card B" not in response.text


@pytest.mark.django_db
class TestKeepPackages:
    def test_the_old_keeper_endpoint_and_page_are_gone(self, logged_in, game, cards):
        """The keep-list folded into packages; its routes must not linger."""
        assert logged_in.post(f"/holdings/keeper/{cards[0].pk}/", {"value": "1"}).status_code == 404
        assert logged_in.get(f"/g/{game.code}/keep-list/").status_code == 404

    def _package_with(self, game, card, superuser, quantity=1):
        set_quantity(card, 3, user=superuser)
        package = SellPackage.objects.create(name="Test lot", game=game, owner=superuser)
        SellPackageItem.objects.create(
            package=package, card=card, quantity=quantity,
            unit_price_at_listing=card.current_price_trend,
        )
        return package

    def test_finalize_decrements_and_audits(self, game, cards, superuser):
        package = self._package_with(game, cards[0], superuser, quantity=2)
        finalize_sale(package, user=superuser, sold_price="5.00", buyer_name="Ana")
        holding = Holding.objects.get(card=cards[0])
        assert holding.quantity == 1
        change = QuantityChange.objects.filter(reason=QuantityChange.Reason.SALE).get()
        assert change.sell_package_id == package.pk
        package.refresh_from_db()
        assert package.status == SellPackage.Status.SOLD
        assert package.buyer_name == "Ana"
        assert package.sold_at is not None

    def test_keepers_are_not_sellable(self, game, cards, superuser):
        package = self._package_with(game, cards[0], superuser, quantity=3)
        reserve(cards[0], superuser, 2)  # only 1 sellable
        with pytest.raises(QuantityError, match="only 1 sellable"):
            finalize_sale(package, user=superuser)
        assert Holding.objects.get(card=cards[0]).quantity == 3  # untouched

    def test_double_finalize_rejected(self, game, cards, superuser):
        package = self._package_with(game, cards[0], superuser)
        finalize_sale(package, user=superuser)
        with pytest.raises(QuantityError, match="already sold"):
            finalize_sale(package, user=superuser)

    def test_empty_package_rejected(self, game, superuser):
        package = SellPackage.objects.create(name="Empty", game=game, owner=superuser)
        with pytest.raises(QuantityError, match="no items"):
            finalize_sale(package, user=superuser)

    def test_package_flow_through_views(self, logged_in, game, cards, superuser):
        set_quantity(cards[0], 2, user=superuser)
        logged_in.post(f"/g/{game.code}/packages/new/", {"name": "Binder lot"})
        package = SellPackage.objects.get()

        logged_in.post(f"/g/{game.code}/packages/{package.pk}/add/",
                       {"card_id": cards[0].pk, "quantity": "2"})
        item = package.items.get()
        assert item.quantity == 2
        assert item.unit_price_at_listing == Decimal("1.00")  # frozen at listing

        # price moves — the frozen listing price must not follow
        Printing.objects.filter(pk=cards[0].pk).update(current_price_trend=Decimal("9.99"))
        item.refresh_from_db()
        assert item.unit_price_at_listing == Decimal("1.00")

        response = logged_in.post(f"/g/{game.code}/packages/{package.pk}/finalize/",
                                  {"sold_price": "3.50"})
        assert response.headers.get("HX-Trigger") == "qty-changed"
        assert Holding.objects.get(card=cards[0]).quantity == 0

    def test_over_committed_package_blocks_finalize_view(self, logged_in, game, cards, superuser):
        set_quantity(cards[0], 1, user=superuser)
        package = SellPackage.objects.create(name="Too big", game=game, owner=superuser)
        SellPackageItem.objects.create(package=package, card=cards[0], quantity=5)
        response = logged_in.post(f"/g/{game.code}/packages/{package.pk}/finalize/", {})
        assert "only 1 sellable" in response.text
        assert Holding.objects.get(card=cards[0]).quantity == 1


@pytest.mark.django_db
class TestExports:
    @pytest.fixture
    def collection(self, cards, superuser):
        set_quantity(cards[0], 2, user=superuser)  # 1.00 each
        set_quantity(cards[1], 1, user=superuser)  # 2.00
        reserve(cards[1], superuser, 1)
        return cards

    def test_xlsx_has_rows_subtotals_and_grand_total(self, game, collection):
        workbook = export_services.build_workbook(game)
        sheet = workbook.active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        assert values[0][0] == "expansion"
        grand = values[-1]
        assert grand[0] == "Grand total"
        assert grand[2] == 3  # 2 + 1 copies
        assert grand[8] == Decimal("4.00")  # 2×1.00 + 1×2.00 on trend

    def test_xlsx_reserved_filters(self, game, collection):
        excluded = export_services.build_workbook(game, keepers="exclude").active
        rows = [[c.value for c in row] for row in excluded.iter_rows()]
        assert rows[-1][2] == 2  # only the non-keeper card's copies
        only = export_services.build_workbook(game, keepers="only").active
        assert [[c.value for c in row] for row in only.iter_rows()][-1][2] == 1

    def test_min_price_filter(self, game, collection):
        sheet = export_services.build_workbook(game, min_price=Decimal("1.50")).active
        assert [[c.value for c in row] for row in sheet.iter_rows()][-1][2] == 1

    def test_xlsx_view_downloads(self, logged_in, game, collection):
        response = logged_in.get(f"/g/{game.code}/export.xlsx")
        assert response.status_code == 200
        assert "attachment" in response.headers["Content-Disposition"]
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active.title == "Collection"

    def test_csv_view_streams_tidy_rows(self, logged_in, game, collection):
        response = logged_in.get(f"/g/{game.code}/export.csv")
        assert response.status_code == 200
        body = b"".join(response.streaming_content).decode()
        lines = [line for line in body.splitlines() if line]
        assert lines[0].startswith("expansion,card_name,quantity")
        assert len(lines) == 3  # header + 2 owned cards, no total row
        assert all(line.split(",")[0] for line in lines[1:])  # no blank continuation cells

    def test_bad_params_fall_back_to_defaults(self, logged_in, game, collection):
        response = logged_in.get(
            f"/g/{game.code}/export.csv?keepers=bogus&basis=bogus&min_price=abc"
        )
        assert response.status_code == 200
        assert len(b"".join(response.streaming_content).decode().strip().splitlines()) == 3
