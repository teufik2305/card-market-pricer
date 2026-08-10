"""Collection exports.

The XLSX reproduces the shape the legacy notebook produced (expansion header
rows with subtotals, indented card rows, grand total), so the spreadsheets you
already keep stay comparable. The CSV is tidy tabular data instead — one row
per card, no merged/blank cells — because that is what a CSV is for.
"""

import csv
from decimal import Decimal

from django.db.models import DecimalField, ExpressionWrapper, F, Sum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from collection.models import Holding, SellPackage, SellPackageItem

BASES = {
    "from": ("current_price_from", "From"),
    "trend": ("current_price_trend", "Trend"),
    "30d": ("current_price_30d", "30-day avg"),
}

HEADERS = [
    "expansion", "card_name", "quantity", "reserved",
    "price_from", "price_trend", "price_30d_avg",
    "total_from", "total_trend", "total_30d",
]


def _line(value, quantity) -> Decimal:
    return (value or Decimal("0")) * quantity


def reserved_card_ids(game):
    """Cards with copies committed to a keep package (the old keeper flag)."""
    return set(
        SellPackageItem.objects.filter(
            package__kind=SellPackage.Kind.KEEP,
            card__expansion__game=game,
        )
        .exclude(package__status=SellPackage.Status.CANCELLED)
        .values_list("card_id", flat=True)
    )


def collection_rows(game, *, keepers: str = "include", min_price=None, basis: str = "trend"):
    """Owned cards of one game, ordered by expansion then card.

    keepers: include | exclude | only — 'keepers' now means copies reserved in a
    keep package, since the standalone keep-list is gone. min_price filters on
    the chosen basis.
    """
    field, _label = BASES.get(basis, BASES["trend"])
    qs = (
        Holding.objects.filter(quantity__gt=0, card__expansion__game=game)
        .select_related("card", "card__expansion")
        .order_by("card__expansion__slug", "card__slug")
    )
    if keepers in ("exclude", "only"):
        reserved = reserved_card_ids(game)
        qs = qs.exclude(card_id__in=reserved) if keepers == "exclude" \
            else qs.filter(card_id__in=reserved)
    if min_price is not None:
        qs = qs.filter(**{f"card__{field}__gte": min_price})
    return qs


def totals_for(rows) -> dict:
    money = lambda expr: ExpressionWrapper(  # noqa: E731
        expr, output_field=DecimalField(max_digits=14, decimal_places=2)
    )
    agg = rows.aggregate(
        copies=Sum("quantity"),
        total_from=Sum(money(F("quantity") * F("card__current_price_from"))),
        total_trend=Sum(money(F("quantity") * F("card__current_price_trend"))),
        total_30d=Sum(money(F("quantity") * F("card__current_price_30d"))),
    )
    return {k: (v if v is not None else Decimal("0")) for k, v in agg.items()}


def _row_values(holding) -> list:
    card = holding.card
    return [
        card.expansion.display_name,
        card.display_name,
        holding.quantity,
        getattr(holding, "reserved", 0),
        card.current_price_from,
        card.current_price_trend,
        card.current_price_30d,
        _line(card.current_price_from, holding.quantity),
        _line(card.current_price_trend, holding.quantity),
        _line(card.current_price_30d, holding.quantity),
    ]


def build_workbook(game, *, keepers="include", min_price=None, basis="trend") -> Workbook:
    """Hierarchical workbook: per-expansion subtotal rows + indented card rows."""
    rows = collection_rows(game, keepers=keepers, min_price=min_price, basis=basis)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Collection"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    group_fill = PatternFill("solid", fgColor="F2F2F2")
    total_fill = PatternFill("solid", fgColor="FFE6B3")
    money_format = '#,##0.00 "€"'

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = header_fill

    current_expansion = None
    group: list = []

    def flush_group():
        if not group:
            return
        subtotal = [
            current_expansion, "", sum(h.quantity for h in group),
            sum(getattr(h, "reserved", 0) for h in group), None, None, None,
            sum(_line(h.card.current_price_from, h.quantity) for h in group),
            sum(_line(h.card.current_price_trend, h.quantity) for h in group),
            sum(_line(h.card.current_price_30d, h.quantity) for h in group),
        ]
        sheet.append(subtotal)
        for cell in sheet[sheet.max_row]:
            cell.font = bold
            cell.fill = group_fill
        for holding in group:
            values = _row_values(holding)
            values[0] = ""  # expansion shown once, on the subtotal row
            values[1] = "    " + values[1]
            sheet.append(values)

    for holding in rows:
        expansion_name = holding.card.expansion.display_name
        if expansion_name != current_expansion:
            flush_group()
            current_expansion, group = expansion_name, []
        group.append(holding)
    flush_group()

    totals = totals_for(rows)
    sheet.append([
        "Grand total", "", totals["copies"], None, None, None, None,
        totals["total_from"], totals["total_trend"], totals["total_30d"],
    ])
    for cell in sheet[sheet.max_row]:
        cell.font = bold
        cell.fill = total_fill

    widths = [38, 46, 10, 10, 12, 12, 12, 13, 13, 13]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=10):
        for cell in row:
            cell.number_format = money_format
            cell.alignment = Alignment(horizontal="right")
    sheet.freeze_panes = "C2"
    return workbook


def csv_lines(game, *, keepers="include", min_price=None, basis="trend"):
    """Generator of CSV lines — tidy rows, streamed, no grand-total row."""

    class Echo:
        def write(self, value):
            return value

    writer = csv.writer(Echo())
    yield writer.writerow(HEADERS)
    for holding in collection_rows(
        game, keepers=keepers, min_price=min_price, basis=basis
    ).iterator(chunk_size=500):
        yield writer.writerow(_row_values(holding))
